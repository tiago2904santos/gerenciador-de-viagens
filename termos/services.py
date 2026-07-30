from __future__ import annotations

import hashlib
import io
import logging

from django.conf import settings

logger = logging.getLogger(__name__)

from cadastros.models import Servidor
from cadastros.models import Viatura
from cadastros.selectors import build_configuracao_context
from core.utils.masks import format_placa

from documentos.services.facade import DocumentoFacade
from documentos.services.facade import DocumentoGerado
from documentos.services.document_cache import build_document_cache_key
from documentos.services.document_cache import build_template_cache_signature
from documentos.services.document_cache import documento_gerado_from_artifact
from documentos.services.document_cache import get_cached_document_artifact
from documentos.services.timing import track_document_generation
from documentos.services.pdf_engine import resolve_pdf_engine
from documentos.services.templates import DocumentTemplateDefinition
from documentos.services.templates import DocumentTemplateRegistry
from documentos.services.templates import default_template_registry
from documentos.services.types import DocumentoFormato
from documentos.services.types import DocumentoTipo
from documentos.services.formatters import format_institucional_rodape_linha

from oficios.documents import VarianteTermo
from oficios.documents import _document_text
from oficios.documents import _format_date
from oficios.documents import _oficio_termo_payload
from oficios.documents import _resolver_variante_padrao
from oficios.documents import _transporte_payload
from oficios.documents import _viagem_payload
from oficios.documents import build_termo_payload
from oficios.models import Oficio

from .models import TermoAutorizacao


_TEMPLATE_DOCX_BY_VARIANTE = {
    VarianteTermo.SEMIPREENCHIDO: "termo_autorizacao.docx",
    VarianteTermo.COMPLETO_COM_VIATURA: "termo_autorizacao_automatico.docx",
    VarianteTermo.COMPLETO_SEM_VIATURA: "termo_autorizacao_automatico_sem_viatura.docx",
}


def listar_servidores_com_termo(oficio: Oficio):
    return oficio.servidores_termo_autorizacao.select_related("cargo", "unidade").order_by("nome")


def preview_termo_context(
    oficio: Oficio,
    servidor: Servidor | None = None,
    *,
    modo_semipreenchido: bool = False,
    variante: str | None = None,
) -> dict:
    servidores_termo = listar_servidores_com_termo(oficio)
    srv = servidor or servidores_termo.first()
    if srv is None:
        return {"erro": "Nenhum servidor selecionado para Termo de Autorizacao neste oficio."}
    if not servidores_termo.filter(pk=srv.pk).exists():
        return {"erro": "Servidor nao selecionado para Termo de Autorizacao neste oficio."}
    payload = build_termo_payload(
        oficio,
        srv,
        modo_semipreenchido=modo_semipreenchido,
        variante=variante,
    )
    return {
        "payload": payload,
        "servidor": srv,
        "servidores": list(servidores_termo),
        "variante_efetiva": payload["termo"]["variante"],
    }


def _facade_termo_com_template(template_docx: str) -> DocumentoFacade:
    registry = DocumentTemplateRegistry()
    for definition in default_template_registry.all():
        if definition.tipo == DocumentoTipo.TERMO_AUTORIZACAO and definition.formato == DocumentoFormato.DOCX:
            definition = DocumentTemplateDefinition(
                tipo=definition.tipo,
                formato=definition.formato,
                template_path=template_docx,
                required_placeholders=definition.required_placeholders,
                stylesheet_paths=definition.stylesheet_paths,
            )
        registry.register(definition)
    return DocumentoFacade(template_registry=registry)


def _legacy_docx_context(payload: dict) -> dict:
    institucional = payload.get("institucional") or {}
    termo = payload.get("termo") or {}
    participante = termo.get("participante") or {}
    viagem = termo.get("viagem") or {}
    transporte = termo.get("transporte") or {}

    endereco_partes = [
        institucional.get("logradouro"),
        institucional.get("numero"),
        institucional.get("bairro"),
        institucional.get("cidade_endereco"),
        institucional.get("uf"),
        institucional.get("cep_formatado"),
    ]
    endereco = ", ".join(str(parte).strip() for parte in endereco_partes if str(parte or "").strip())

    return {
        "unidade": institucional.get("unidade") or termo.get("oficio", {}).get("origem") or "",
        "divisao": institucional.get("divisao") or "",
        "endereco": endereco,
        "telefone": participante.get("telefone_formatado") or "",
        "email": institucional.get("email") or "",
        "unidade_rodape": format_institucional_rodape_linha(institucional),
        "data_do_evento": viagem.get("periodo") or viagem.get("saida") or "",
        "destino": viagem.get("destinos_texto") or viagem.get("destino_principal") or "",
        "nome_servidor": participante.get("nome") or "",
        "rg_servidor": participante.get("rg_formatado") or "",
        "cpf_servidor": participante.get("cpf_formatado") or participante.get("cpf") or "",
        "lotacao": participante.get("unidade") or "",
        "viatura": transporte.get("modelo") or "",
        "placa": transporte.get("placa") or "",
        "combustivel": transporte.get("combustivel") or "",
    }


@track_document_generation("termo_gerar_documento")
def gerar_termo_um(
    oficio: Oficio,
    servidor: Servidor,
    formato: DocumentoFormato,
    *,
    modo_semipreenchido: bool = False,
    variante: str | None = None,
) -> DocumentoGerado:
    payload = build_termo_payload(
        oficio,
        servidor,
        modo_semipreenchido=modo_semipreenchido,
        variante=variante,
    )
    ref = f"{oficio.numero_formatado.replace('/', '-')}-termo-{servidor.pk}"
    variante_efetiva = payload["termo"]["variante"]
    template_docx = _TEMPLATE_DOCX_BY_VARIANTE.get(variante_efetiva, "termo_autorizacao.docx")
    docxtpl_context = _legacy_docx_context(payload)
    attempt_chain = ()
    if formato == DocumentoFormato.PDF:
        attempt_chain = resolve_pdf_engine(
            explicit_setting=getattr(settings, "DOCUMENTOS_DEFAULT_PDF_ENGINE", "auto"),
            prefer_docx_pipeline=True,
        ).attempt_chain
    cache_key = build_document_cache_key(
        tipo=DocumentoTipo.TERMO_AUTORIZACAO,
        formato=formato,
        reference=ref,
        payload=payload,
        docxtpl_context=docxtpl_context,
        attempt_chain=attempt_chain,
        template_signature=build_template_cache_signature(
            tipo=DocumentoTipo.TERMO_AUTORIZACAO,
            formato=formato,
            docx_template_path=template_docx,
        ),
    )
    cached = get_cached_document_artifact(
        oficio_id=oficio.pk,
        servidor_id=servidor.pk,
        tipo=DocumentoTipo.TERMO_AUTORIZACAO,
        formato=formato,
        cache_key=cache_key,
    )
    if cached is not None:
        return documento_gerado_from_artifact(
            cached,
            tipo=DocumentoTipo.TERMO_AUTORIZACAO,
            formato=formato,
            reference=ref,
        )

    facade = _facade_termo_com_template(template_docx)
    doc = facade.gerar(
        tipo=DocumentoTipo.TERMO_AUTORIZACAO,
        formato=formato,
        payload=payload,
        reference=ref,
        docxtpl_context=docxtpl_context,
    )
    _persistir_termo_artefato(
        oficio,
        servidor,
        doc,
        cache_key=cache_key,
        payload_snapshot=payload,
    )
    return doc


def _persistir_termo_artefato(
    oficio: Oficio,
    servidor: Servidor,
    doc: DocumentoGerado,
    *,
    cache_key: str = "",
    payload_snapshot: dict | None = None,
) -> None:
    """Persiste o termo como ``DocumentoArtefato`` (para auto-organização no Drive).

    Idempotente por conteúdo (hash). É no-op quando a persistência está desligada
    (ex.: testes). Falhas nunca quebram a geração do documento.
    """
    try:
        from documentos.models import DocumentoArtefato
        from documentos.services.persistence import persist_geracao
        from integracoes.google_drive import naming

        if DocumentoArtefato.objects.filter(
            tipo=DocumentoTipo.TERMO_AUTORIZACAO.value, hash_sha256=doc.hash_sha256
        ).exists():
            return
        cidade = naming.cidade_evento(getattr(oficio, "evento", None), oficio)
        persist_geracao(
            doc,
            oficio_id=oficio.pk,
            servidor_id=servidor.pk,
            evento_id=getattr(oficio, "evento_id", None),
            nome_drive=naming.nome_termo(oficio, servidor, cidade),
            payload_snapshot=payload_snapshot,
            cache_key=cache_key,
            engine=doc.pdf_engine_used or "",
        )
    except Exception:
        logger.warning("Não foi possível persistir artefato do termo.", exc_info=True)


def gerar_termo_lote(oficio: Oficio, formato: DocumentoFormato) -> list[DocumentoGerado]:
    out: list[DocumentoGerado] = []
    for servidor in listar_servidores_com_termo(oficio):
        out.append(gerar_termo_um(oficio, servidor, formato))
    return out


def resolver_artefato_termo_oficio(oficio: Oficio, servidor: Servidor):
    """Garante que exista um `DocumentoArtefato` PDF do termo embutido no ofício.

    Reaproveita o mais recente se já existir; senão gera e persiste agora. Usado
    para oferecer "anexar assinado" mesmo que ninguém tenha aberto o preview antes.
    """
    from documentos.selectors import get_latest_artefato_pdf_termo

    artefato = get_latest_artefato_pdf_termo(oficio.pk, servidor.pk)
    if artefato is not None:
        return artefato
    gerar_termo_um(oficio, servidor, DocumentoFormato.PDF)
    return get_latest_artefato_pdf_termo(oficio.pk, servidor.pk)


def _participante_payload(servidor: Servidor | None) -> dict:
    if servidor is None:
        return {
            "id": "",
            "nome": "",
            "cargo": "",
            "rg_formatado": "",
            "cpf": "",
            "cpf_formatado": "",
            "unidade": "",
            "telefone_formatado": "",
            "email": "",
        }
    return {
        "id": servidor.pk,
        "nome": _document_text(servidor.nome),
        "cargo": _document_text(servidor.cargo.nome if servidor.cargo_id else ""),
        "rg_formatado": servidor.rg_formatado,
        "cpf": servidor.cpf or "",
        "cpf_formatado": servidor.cpf_formatado,
        "unidade": _document_text(servidor.unidade.nome if servidor.unidade_id else ""),
        "telefone_formatado": servidor.telefone_formatado,
        "email": "",
    }


def _viatura_payload(viatura: Viatura | None) -> dict:
    if viatura is None:
        return {
            "tem_viatura": False,
            "placa": "-",
            "modelo": "-",
            "tipo": "-",
            "combustivel": "-",
            "motorista": "-",
            "porte_armas": "Sim",
            "unidade": "-",
        }
    return {
        "tem_viatura": True,
        "placa": format_placa(viatura.placa),
        "modelo": _document_text(viatura.modelo),
        "tipo": _document_text(viatura.get_tipo_display() if viatura.tipo else ""),
        "combustivel": _document_text(str(viatura.combustivel) if viatura.combustivel_id else ""),
        "motorista": "-",
        "porte_armas": "Sim",
        "unidade": _document_text(viatura.unidade if viatura.unidade_id else ""),
    }


_MESES = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril",
    5: "maio", 6: "junho", 7: "julho", 8: "agosto",
    9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro",
}


def _fmt_extenso(d) -> str:
    return f"{d.day} de {_MESES[d.month]} de {d.year}"


def _periodo_texto(termo: TermoAutorizacao) -> str:
    inicio, fim = termo.periodo_efetivo()
    if not inicio:
        return "-"
    if not fim or fim == inicio:
        return f"no dia {_fmt_extenso(inicio)}"
    if inicio.month == fim.month and inicio.year == fim.year:
        return f"nos dias {inicio.day} até {fim.day} de {_MESES[inicio.month]} de {inicio.year}"
    if inicio.year == fim.year:
        return (
            f"nos dias {inicio.day} de {_MESES[inicio.month]} até "
            f"{fim.day} de {_MESES[fim.month]} de {fim.year}"
        )
    return f"nos dias {_fmt_extenso(inicio)} até {_fmt_extenso(fim)}"


def _viagem_payload_termo(termo: TermoAutorizacao) -> dict:
    if termo.oficio_id and not termo.destino_cidade_id and not termo.data_evento_inicio:
        payload = _viagem_payload(termo.oficio)
        payload["periodo"] = _periodo_texto(termo)
        return payload

    destino = termo.destino_efetivo() or "-"
    inicio, fim = termo.periodo_efetivo()
    saida = inicio.strftime("%d/%m/%Y") if inicio else "-"
    retorno = fim.strftime("%d/%m/%Y") if fim else saida
    return {
        "destino_principal": destino,
        "destinos_texto": destino,
        "saida": saida,
        "retorno": retorno,
        "periodo": _periodo_texto(termo),
        "roteiro_ida": [destino] if destino != "-" else [],
        "roteiro_ida_texto": destino,
        "roteiro_retorno": [],
        "roteiro_retorno_texto": "-",
        "quantidade_diarias": "-",
        "valor_diarias": "-",
        "valor_diarias_extenso": "-",
        "motivo": "Termo de autorizacao para deslocamento em evento.",
    }


def _oficio_payload_termo(termo: TermoAutorizacao) -> dict:
    if termo.oficio_id:
        return _oficio_termo_payload(termo.oficio)
    return {
        "numero_formatado": f"Termo #{termo.pk or 'novo'}",
        "protocolo_formatado": "-",
        "data_criacao": _format_date(termo.created_at.date() if termo.created_at else None),
        "assunto": "Termo de autorizacao avulso",
        "origem": build_configuracao_context(area=getattr(termo, "area", None)).get("unidade") or "-",
        "destino": "Direcao/chefia competente",
    }


def _transporte_payload_termo(termo: TermoAutorizacao) -> dict:
    if termo.viatura_id:
        return _viatura_payload(termo.viatura)
    if termo.oficio_id:
        return _transporte_payload(termo.oficio)
    return _viatura_payload(None)


def _resolver_variante_termo_cadastro(termo: TermoAutorizacao, servidor: Servidor | None) -> str:
    if servidor is None:
        return VarianteTermo.SEMIPREENCHIDO
    if termo.viatura_id:
        return VarianteTermo.COMPLETO_COM_VIATURA
    if termo.oficio_id:
        return _resolver_variante_padrao(termo.oficio)
    return VarianteTermo.COMPLETO_SEM_VIATURA


def build_termo_cadastro_payload(
    termo: TermoAutorizacao,
    servidor: Servidor | None = None,
) -> dict:
    area = getattr(termo, "area", None) or getattr(getattr(termo, "oficio", None), "area", None)
    institucional = build_configuracao_context(area=area)
    payload = {
        "institucional": institucional,
        "oficio": _oficio_payload_termo(termo),
    }
    variante = _resolver_variante_termo_cadastro(termo, servidor)
    payload["termo"] = {
        "variante": variante,
        "participante": _participante_payload(servidor),
        "viagem": _viagem_payload_termo(termo),
        "transporte": _transporte_payload_termo(termo),
        "oficio": _oficio_payload_termo(termo),
        "textos": {
            "titulo": "TERMO DE AUTORIZACAO",
            "corpo_autorizacao": (
                "Autorizo o servidor acima identificado a realizar o deslocamento descrito neste "
                "termo, observadas as normas administrativas aplicaveis ao servico publico e ao "
                "uso do transporte informado."
            ),
            "declaracao": (
                "O servidor declara ciencia das informacoes registradas, do periodo autorizado "
                "e das responsabilidades funcionais decorrentes da viagem."
            ),
            "observacoes": "-",
        },
    }
    return payload


def servidores_para_termo_cadastro(termo: TermoAutorizacao) -> list[Servidor | None]:
    servidores = list(termo.servidores_efetivos())
    return servidores or [None]


def gerar_termo_cadastro_um(
    termo: TermoAutorizacao,
    servidor: Servidor | None,
    formato: DocumentoFormato,
) -> DocumentoGerado:
    payload = build_termo_cadastro_payload(termo, servidor)
    variante_efetiva = payload["termo"]["variante"]
    template_docx = _TEMPLATE_DOCX_BY_VARIANTE.get(variante_efetiva, "termo_autorizacao.docx")
    facade = _facade_termo_com_template(template_docx)
    ref_servidor = servidor.pk if servidor is not None else "sem-servidor"
    ref = f"termo-{termo.pk}-cadastro-{ref_servidor}"
    doc = facade.gerar(
        tipo=DocumentoTipo.TERMO_AUTORIZACAO,
        formato=formato,
        payload=payload,
        reference=ref,
        docxtpl_context=_legacy_docx_context(payload),
    )
    if termo.pk:
        _persistir_termo_cadastro_artefato(termo, servidor, doc)
    return doc


def _persistir_termo_cadastro_artefato(
    termo: TermoAutorizacao, servidor: Servidor | None, doc: DocumentoGerado
) -> None:
    """Persiste o termo avulso/cadastro como ``DocumentoArtefato`` (Drive + anexar assinado).

    Análogo a ``_persistir_termo_artefato``, mas para o cadastro independente
    (``TermoAutorizacao``), que pode não ter ofício vinculado — por isso usa
    ``termo_id`` (em vez de só ``oficio_id``) para desambiguar termos avulsos
    que compartilhem o mesmo servidor (ou sejam ambos genéricos, sem servidor).
    """
    try:
        from documentos.models import DocumentoArtefato
        from documentos.services.persistence import persist_geracao
        from integracoes.google_drive import naming

        servidor_id = servidor.pk if servidor is not None else None
        if DocumentoArtefato.objects.filter(
            tipo=DocumentoTipo.TERMO_AUTORIZACAO.value,
            termo_id=termo.pk,
            servidor_id=servidor_id,
            hash_sha256=doc.hash_sha256,
        ).exists():
            return
        oficio = termo.oficio
        nome_drive = ""
        if oficio is not None:
            try:
                cidade = naming.cidade_evento(getattr(oficio, "evento", None), oficio)
                nome_drive = naming.nome_termo(oficio, servidor, cidade)
            except Exception:
                nome_drive = ""
        persist_geracao(
            doc,
            oficio_id=oficio.pk if oficio is not None else None,
            servidor_id=servidor_id,
            evento_id=termo.evento_id,
            termo_id=termo.pk,
            nome_drive=nome_drive,
        )
    except Exception:
        logger.warning("Não foi possível persistir artefato do termo (cadastro).", exc_info=True)


def gerar_termo_cadastro_lote(termo: TermoAutorizacao, formato: DocumentoFormato) -> list[DocumentoGerado]:
    return [gerar_termo_cadastro_um(termo, servidor, formato) for servidor in servidores_para_termo_cadastro(termo)]


def resolver_artefato_termo_cadastro(termo: TermoAutorizacao, servidor: Servidor | None):
    """Garante que exista um `DocumentoArtefato` PDF do termo avulso (genérico ou por servidor).

    Reaproveita o mais recente se já existir; senão gera e persiste agora.
    """
    from documentos.selectors import get_latest_artefato_pdf_termo_cadastro

    servidor_id = servidor.pk if servidor is not None else None
    artefato = get_latest_artefato_pdf_termo_cadastro(termo.pk, servidor_id)
    if artefato is not None:
        return artefato
    gerar_termo_cadastro_um(termo, servidor, DocumentoFormato.PDF)
    return get_latest_artefato_pdf_termo_cadastro(termo.pk, servidor_id)


def termo_oficio_tem_assinado(oficio: Oficio, servidor: Servidor) -> bool:
    """``True`` se já existe uma versão assinada anexada para este termo embutido no ofício."""
    from documentos.selectors import get_latest_artefato_pdf_termo

    artefato = get_latest_artefato_pdf_termo(oficio.pk, servidor.pk)
    return artefato is not None and artefato.esta_assinado


def pdf_termo_oficio_assinado_ou_gerado(oficio: Oficio, servidor: Servidor) -> bytes:
    """Bytes do termo embutido no ofício: prefere o assinado anexado, senão gera."""
    from documentos.selectors import get_latest_artefato_pdf_termo

    artefato = get_latest_artefato_pdf_termo(oficio.pk, servidor.pk)
    if artefato is not None and artefato.esta_assinado:
        with artefato.arquivo_assinado.open("rb") as f:
            return f.read()
    return gerar_termo_um(oficio, servidor, DocumentoFormato.PDF).conteudo


def termo_oficio_assinado_info(oficio: Oficio, servidor: Servidor) -> dict:
    """Dados de assinatura de um termo embutido no ofício, para presenters de card/wizard.

    Retorna ``assinado``, ``anexar_assinado_url``, ``assinado_nome_original``,
    ``assinado_view_url`` e ``remover_assinado_url``. Se ainda não existe artefato
    persistido, ``anexar_assinado_url`` aponta para o wrapper que resolve-ou-gera
    sob demanda (``termos:termo_oficio_assinado_anexar``).
    """
    from django.urls import reverse

    from documentos.selectors import get_latest_artefato_pdf_termo

    artefato = get_latest_artefato_pdf_termo(oficio.pk, servidor.pk)
    if artefato is not None:
        assinado = artefato.esta_assinado
        return {
            "assinado": assinado,
            "anexar_assinado_url": reverse("documentos:artefato_assinado_anexar", args=[artefato.pk]),
            "assinado_nome_original": artefato.assinado_nome_original if assinado else "",
            "assinado_view_url": (
                reverse("documentos:artefato_pdf_conteudo", args=[artefato.pk]) if assinado else ""
            ),
            "remover_assinado_url": (
                reverse("documentos:artefato_assinado_remover", args=[artefato.pk]) if assinado else ""
            ),
        }
    return {
        "assinado": False,
        "anexar_assinado_url": reverse("termos:termo_oficio_assinado_anexar", args=[oficio.pk, servidor.pk]),
        "assinado_nome_original": "",
        "assinado_view_url": "",
        "remover_assinado_url": "",
    }


def termo_cadastro_tem_assinado(termo: TermoAutorizacao, servidor: Servidor | None) -> bool:
    """``True`` se já existe uma versão assinada anexada para este termo avulso (genérico ou por servidor)."""
    from documentos.selectors import get_latest_artefato_pdf_termo_cadastro

    servidor_id = servidor.pk if servidor is not None else None
    artefato = get_latest_artefato_pdf_termo_cadastro(termo.pk, servidor_id)
    return artefato is not None and artefato.esta_assinado


def pdf_termo_cadastro_assinado_ou_gerado(termo: TermoAutorizacao, servidor: Servidor | None) -> bytes:
    """Bytes do termo avulso/cadastro: prefere o assinado anexado, senão gera."""
    from documentos.selectors import get_latest_artefato_pdf_termo_cadastro

    servidor_id = servidor.pk if servidor is not None else None
    artefato = get_latest_artefato_pdf_termo_cadastro(termo.pk, servidor_id)
    if artefato is not None and artefato.esta_assinado:
        with artefato.arquivo_assinado.open("rb") as f:
            return f.read()
    return gerar_termo_cadastro_um(termo, servidor, DocumentoFormato.PDF).conteudo


def termo_cadastro_assinado_info(termo: TermoAutorizacao, servidor_id: int | None) -> dict:
    """Dados de assinatura de um termo avulso (genérico ou por servidor) para presenters de lista.

    Retorna ``assinado``, ``anexar_assinado_url``, ``assinado_nome_original`` e
    ``assinado_view_url`` prontos para uso em cards/linhas de lista e no modal
    "anexar assinado" (termos/views.py e eventos/views.py).
    """
    from django.urls import reverse

    from documentos.selectors import get_latest_artefato_pdf_termo_cadastro

    artefato = get_latest_artefato_pdf_termo_cadastro(termo.pk, servidor_id)
    if artefato is not None:
        assinado = artefato.esta_assinado
        return {
            "assinado": assinado,
            "anexar_assinado_url": reverse("documentos:artefato_assinado_anexar", args=[artefato.pk]),
            "assinado_nome_original": artefato.assinado_nome_original if assinado else "",
            "assinado_view_url": (
                reverse("documentos:artefato_pdf_conteudo", args=[artefato.pk]) if assinado else ""
            ),
            "remover_assinado_url": (
                reverse("documentos:artefato_assinado_remover", args=[artefato.pk]) if assinado else ""
            ),
        }
    if servidor_id is None:
        anexar_url = reverse("termos:termo_cadastro_generico_assinado_anexar", args=[termo.pk])
    else:
        anexar_url = reverse("termos:termo_cadastro_servidor_assinado_anexar", args=[termo.pk, servidor_id])
    return {
        "assinado": False,
        "anexar_assinado_url": anexar_url,
        "assinado_nome_original": "",
        "assinado_view_url": "",
        "remover_assinado_url": "",
    }


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def fundir_termos_docx(documentos: list[DocumentoGerado]) -> bytes:
    from docx import Document as DocxDocument
    from docxcompose.composer import Composer

    base = DocxDocument(io.BytesIO(documentos[0].conteudo))
    composer = Composer(base)
    for doc in documentos[1:]:
        base.add_page_break()
        composer.append(DocxDocument(io.BytesIO(doc.conteudo)))
    buf = io.BytesIO()
    composer.save(buf)
    return buf.getvalue()


def fundir_termos_pdf(documentos: list[DocumentoGerado]) -> bytes:
    return fundir_termos_pdf_bytes([doc.conteudo for doc in documentos])


def gerar_termos_pdf_consolidado(oficio: Oficio) -> bytes:
    """Gera o lote com uma única conversão DOCX→PDF."""
    documentos = gerar_termo_lote(oficio, DocumentoFormato.DOCX)
    if not documentos:
        return b""
    docx = documentos[0].conteudo if len(documentos) == 1 else fundir_termos_docx(documentos)
    facade = DocumentoFacade()
    pdf, _engine = facade.converter_docx_pronto_para_pdf(
        docx,
        tipo=DocumentoTipo.TERMO_AUTORIZACAO,
    )
    return pdf


def fundir_termos_pdf_bytes(conteudos: list[bytes]) -> bytes:
    """Como ``fundir_termos_pdf``, mas a partir de bytes já resolvidos (assinado ou gerado)."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for conteudo in conteudos:
        reader = PdfReader(io.BytesIO(conteudo))
        for page in reader.pages:
            writer.add_page(page)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()
